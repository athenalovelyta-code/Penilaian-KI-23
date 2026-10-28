import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(master_data, tabmin_data, product_data, profesi_data) -> io.BytesIO:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling palettes
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark slate
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    sub_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    sub_header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    
    stripe_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    # ==========================================
    # SHEET 1: REKAPITULASI NILAI AKHIR
    # ==========================================
    ws1 = wb.create_sheet(title="Rekap Nilai Akhir")
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws1.merge_cells("A1:K1")
    ws1["A1"] = "SISTEM PENILAIAN ACARA KI23 - REKAPITULASI NILAI AKHIR"
    ws1["A1"].font = Font(name="Arial", size=14, bold=True, color="1E293B")
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.row_dimensions[1].height = 28

    headers1 = [
        "No", "Nama Peserta", "Divisi Profesi", "Kelompok Tabmin", "PJK",
        "Nilai Kelas Profesi (1-5)", "Nilai Tabmin Anggota (1-5)", "Nilai Produk Kelompok (1-5)",
        "Nilai Akhir (Skala 1-5)", "Nilai Akhir (Skala 100)", "Predikat"
    ]
    ws1.append([]) # Row 2 empty
    ws1.append(headers1) # Row 3
    ws1.row_dimensions[3].height = 26

    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(3, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    row_num = 4
    for idx, item in enumerate(master_data, start=1):
        row_vals = [
            idx,
            item["name"],
            item["profession"],
            item["tabmin_group"],
            item["pjk_name"],
            item["nilai_profesi"],
            item["nilai_tabmin_anggota"],
            item["nilai_produk_kelompok"],
            item["nilai_akhir_skala_5"],
            item["nilai_akhir_skala_100"],
            item["predikat"]
        ]
        ws1.append(row_vals)
        current_fill = stripe_fill if idx % 2 == 0 else white_fill
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws1.cell(row_num, col_idx)
            cell.fill = current_fill
            cell.border = thin_border
            if col_idx in [1, 3, 4, 5, 11]:
                cell.alignment = center_align
            elif col_idx == 2:
                cell.alignment = left_align
            else:
                cell.alignment = right_align
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
        row_num += 1

    # Auto-adjust column widths
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ==========================================
    # SHEET 2: DETAIL TABMIN ANGGOTA
    # ==========================================
    ws2 = wb.create_sheet(title="Detail Tabmin Anggota")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:J1")
    ws2["A1"] = "DETAIL PENILAIAN ANGGOTA TABMIN (HRD vs PJK vs PEER)"
    ws2["A1"].font = Font(name="Arial", size=14, bold=True, color="1E293B")
    ws2.row_dimensions[1].height = 28

    headers2 = [
        "No", "Nama Peserta", "Kelompok", "PJK",
        "Skor HRD", "Skor PJK", "Rata-rata Peer", "Jml Reviewer Peer",
        "Nilai Akhir Anggota (1-5)", "Status Penilai"
    ]
    ws2.append([])
    ws2.append(headers2)
    ws2.row_dimensions[3].height = 26

    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(3, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    row_num = 4
    for idx, item in enumerate(tabmin_data, start=1):
        row_vals = [
            idx,
            item["name"],
            item["tabmin_group"],
            item["pjk_name"],
            item["skor_hrd"],
            item["skor_pjk"],
            item["skor_peer"],
            item["peer_count"],
            item["nilai_tabmin_anggota"],
            item["status_penilai"]
        ]
        ws2.append(row_vals)
        current_fill = stripe_fill if idx % 2 == 0 else white_fill
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws2.cell(row_num, col_idx)
            cell.fill = current_fill
            cell.border = thin_border
            if col_idx in [1, 3, 4, 8, 10]:
                cell.alignment = center_align
            elif col_idx == 2:
                cell.alignment = left_align
            else:
                cell.alignment = right_align
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
        row_num += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ==========================================
    # SHEET 3: DETAIL PRODUK TABMIN
    # ==========================================
    ws3 = wb.create_sheet(title="Detail Produk Tabmin")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.merge_cells("A1:G1")
    ws3["A1"] = "PENILAIAN PRODUK / HASIL KARYA KELOMPOK TABMIN"
    ws3["A1"].font = Font(name="Arial", size=14, bold=True, color="1E293B")
    ws3.row_dimensions[1].height = 28

    headers3 = [
        "Kelompok", "Hasil Reporter", "Hasil Fotografer", "Hasil Ilustrator",
        "Hasil Layouter", "Hasil Marcomm", "Penilaian Total Produk"
    ]
    ws3.append([])
    ws3.append(headers3)
    ws3.row_dimensions[3].height = 26

    for col_idx in range(1, len(headers3) + 1):
        cell = ws3.cell(3, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    row_num = 4
    for idx, item in enumerate(product_data, start=1):
        row_vals = [
            item["tabmin_group"],
            item["score_reporter"],
            item["score_fotografer"],
            item["score_ilustrator"],
            item["score_layouter"],
            item["score_marcomm"],
            item["score_total"]
        ]
        ws3.append(row_vals)
        current_fill = stripe_fill if idx % 2 == 0 else white_fill
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws3.cell(row_num, col_idx)
            cell.fill = current_fill
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = center_align
            else:
                cell.alignment = right_align
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
        row_num += 1

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # ==========================================
    # SHEET 4: DETAIL KELAS PROFESI
    # ==========================================
    ws4 = wb.create_sheet(title="Detail Kelas Profesi")
    ws4.views.sheetView[0].showGridLines = True
    
    ws4.merge_cells("A1:K1")
    ws4["A1"] = "DETAIL PENILAIAN KELAS PROFESI (KEPROF 1, 2, 3)"
    ws4["A1"].font = Font(name="Arial", size=14, bold=True, color="1E293B")
    ws4.row_dimensions[1].height = 28

    headers4 = [
        "No", "Nama Peserta", "Profesi", "Kelompok",
        "Keprof 1 (Umum / Tugas Pengganti)", "Keprof 1 (Tugas Utama)",
        "Keprof 2 (Umum / Tugas Pengganti)", "Keprof 2 (Tugas Utama)",
        "Keprof 3 (Umum)", "Rata-rata Kelas Profesi", "Status Absensi / Pengganti"
    ]
    ws4.append([])
    ws4.append(headers4)
    ws4.row_dimensions[3].height = 26

    for col_idx in range(1, len(headers4) + 1):
        cell = ws4.cell(3, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    row_num = 4
    for idx, item in enumerate(profesi_data, start=1):
        row_vals = [
            idx,
            item["name"],
            item["profession"],
            item["tabmin_group"],
            item["keprof1_umum"],
            item["keprof1_tugas"],
            item["keprof2_umum"],
            item["keprof2_tugas"],
            item["keprof3_umum"],
            item["nilai_profesi_total"],
            item["status_absen"]
        ]
        ws4.append(row_vals)
        current_fill = stripe_fill if idx % 2 == 0 else white_fill
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws4.cell(row_num, col_idx)
            cell.fill = current_fill
            cell.border = thin_border
            if col_idx in [1, 3, 4, 11]:
                cell.alignment = center_align
            elif col_idx == 2:
                cell.alignment = left_align
            else:
                cell.alignment = right_align
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
        row_num += 1

    for col in ws4.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws4.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
