import argparse
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

from server.database import get_db


ROOT_DIR = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT_DIR / "ABSENSI KEPROF.xlsx"
ATTENDANCE_BLOCKS = (
    (0, "Ilustrator"),
    (4, "Fotografer"),
    (8, "Layouter"),
    (12, "Marcomm"),
    (16, "Reporter"),
)
IMPORT_EVALUATOR_NAME = "Impor ABSENSI KEPROF.xlsx"


def normalize_name(value):
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    without_accents = "".join(
        char for char in normalized if not unicodedata.combining(char)
    )
    return re.sub(r"[^a-z0-9]+", "", without_accents.casefold())


def is_checked(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1
    return str(value or "").strip().casefold() in {"true", "1", "yes", "y", "✓", "✔"}


def read_attendance(workbook_path, participants, duplicate_resolutions=None):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook.active
    participants_by_name = {}
    duplicate_resolutions = duplicate_resolutions or {}
    for participant in participants:
        key = normalize_name(participant["name"])
        participants_by_name.setdefault(key, []).append(participant)

    attendance_rows = []
    seen_participant_ids = set()
    errors = []

    try:
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=5, max_col=20, values_only=True),
            start=5,
        ):
            for name_column, expected_profession in ATTENDANCE_BLOCKS:
                name = row[name_column]
                if not name or not str(name).strip():
                    continue

                matching_participants = participants_by_name.get(normalize_name(name), [])
                if not matching_participants:
                    errors.append(f"Baris {row_number}: peserta '{name}' tidak ditemukan.")
                    continue
                if len(matching_participants) > 1:
                    requested_group = duplicate_resolutions.get(normalize_name(name))
                    matching_participants = [
                        participant for participant in matching_participants
                        if participant["tabmin_group"] == requested_group
                    ]
                    if len(matching_participants) != 1:
                        choices = ", ".join(
                            f"{participant['name']} ({participant['tabmin_group']})"
                            for participant in participants_by_name[normalize_name(name)]
                        )
                        errors.append(
                            f"Baris {row_number}: peserta '{name}' ambigu ({choices}); "
                            "tentukan kelompok yang benar."
                        )
                        continue
                participant = matching_participants[0]
                if participant["id"] in seen_participant_ids:
                    errors.append(f"Baris {row_number}: peserta '{name}' tercatat lebih dari sekali.")
                    continue
                if participant["profession"].casefold() != expected_profession.casefold():
                    errors.append(
                        f"Baris {row_number}: profesi '{name}' pada file adalah "
                        f"{expected_profession}, tetapi di database "
                        f"{participant['profession']}."
                    )
                    continue

                seen_participant_ids.add(participant["id"])
                for session, value in enumerate(row[name_column + 1:name_column + 4], start=1):
                    attendance_rows.append(
                        (
                            participant["id"],
                            session,
                            0 if is_checked(value) else 1,
                            IMPORT_EVALUATOR_NAME,
                        )
                    )
    finally:
        workbook.close()

    if errors:
        raise ValueError(
            "Impor dibatalkan karena ada data absensi yang tidak cocok:\n"
            + "\n".join(errors)
        )
    if not attendance_rows:
        raise ValueError("Tidak ada data absensi peserta yang ditemukan di workbook.")

    return attendance_rows, seen_participant_ids


def import_attendance(workbook_path, dry_run=False, duplicate_resolutions=None):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, name, profession, tabmin_group FROM participants")
        participants = cursor.fetchall()
        attendance_rows, participant_ids = read_attendance(
            workbook_path,
            participants,
            duplicate_resolutions=duplicate_resolutions,
        )

        if not dry_run:
            cursor.executemany(
                """
                INSERT INTO profesi_attendance
                    (participant_id, keprof_session, is_absen, evaluator_id,
                     evaluator_name, updated_at)
                VALUES (?, ?, ?, NULL, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(participant_id, keprof_session) DO UPDATE SET
                    is_absen = excluded.is_absen,
                    evaluator_id = NULL,
                    evaluator_name = excluded.evaluator_name,
                    updated_at = CURRENT_TIMESTAMP
                """,
                attendance_rows,
            )

    present_count = sum(1 for row in attendance_rows if row[2] == 0)
    absent_count = len(attendance_rows) - present_count
    print(
        f"{'Simulasi' if dry_run else 'Impor selesai'}: "
        f"{len(participant_ids)} peserta, {len(attendance_rows)} status sesi, "
        f"{present_count} hadir, {absent_count} tidak hadir."
    )
    return len(attendance_rows)


def main():
    parser = argparse.ArgumentParser(
        description="Impor absensi Keprof dari workbook ke database aplikasi."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Lokasi file ABSENSI KEPROF.xlsx.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validasi data dan tampilkan ringkasan tanpa mengubah database.",
    )
    parser.add_argument(
        "--resolve-duplicate",
        action="append",
        default=[],
        metavar="NAME=GROUP",
        help="Pilih kelompok untuk nama peserta duplikat, misalnya 'Shofiyyah Mukhayadi=Kelompok 6'.",
    )
    args = parser.parse_args()

    if not args.workbook.is_file():
        parser.error(f"File absensi tidak ditemukan: {args.workbook}")

    duplicate_resolutions = {}
    for resolution in args.resolve_duplicate:
        if "=" not in resolution:
            parser.error("--resolve-duplicate harus berformat NAME=GROUP.")
        name, group = resolution.rsplit("=", 1)
        duplicate_resolutions[normalize_name(name)] = group.strip()

    import_attendance(
        args.workbook,
        dry_run=args.dry_run,
        duplicate_resolutions=duplicate_resolutions,
    )


if __name__ == "__main__":
    main()
